from pathlib import Path
import sys

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DARK = '0F172A'
ACCENT = '1D4ED8'
WHITE = 'FFFFFF'


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = pr.find(qn('w:shd'))
    if el is None:
        el = OxmlElement('w:shd')
        pr.append(el)
    el.set(qn('w:fill'), fill)


def border(cell, color=DARK, size='10'):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        el = borders.find(qn('w:' + edge))
        if el is None:
            el = OxmlElement('w:' + edge)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)


def margins(cell, top=55, start=75, bottom=55, end=75):
    pr = cell._tc.get_or_add_tcPr()
    m = pr.first_child_found_in('w:tcMar')
    if m is None:
        m = OxmlElement('w:tcMar')
        pr.append(m)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        el = m.find(qn('w:' + name))
        if el is None:
            el = OxmlElement('w:' + name)
            m.append(el)
        el.set(qn('w:w'), str(value))
        el.set(qn('w:type'), 'dxa')


def fixed_widths(table, values):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    pr = table._tbl.tblPr
    layout = pr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        pr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    for i, width in enumerate(values):
        table.columns[i].width = Cm(width)
        for row in table.rows:
            row.cells[i].width = Cm(width)


def clear(cell):
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    return p


def section_elements(doc, start_prefix, end_prefix):
    body = list(doc.element.body)
    start = end = None
    for i, el in enumerate(body):
        if el.tag != qn('w:p'):
            continue
        text = Paragraph(el, doc).text.strip()
        if start is None and text.startswith(start_prefix):
            start = i
        elif start is not None and text.startswith(end_prefix):
            end = i
            break
    if start is None or end is None:
        raise RuntimeError(f'Missing section: {start_prefix}')
    return body[start + 1:end]


def style_run(run, *, bold=False, color=DARK, size=10):
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def rebuild_altgr(doc):
    elems = section_elements(doc, '3. Sonderzeichen mit AltGr', '4. Programme & Browser')
    instruction = next(
        Paragraph(el, doc)
        for el in elems
        if el.tag == qn('w:p') and Paragraph(el, doc).text.strip()
    )
    old_table = next(Table(el, doc) for el in elems if el.tag == qn('w:tbl'))

    instruction.text = (
        'Schreibe in jedes stark umrandete Feld nur die zweite Taste. '
        'Lies jede Aufgabe als: ZEICHEN = AltGr + [2. Taste]. '
        'Es gilt die Belegung einer Schweizer Tastatur.'
    )
    instruction.paragraph_format.space_after = Pt(4)
    for run in instruction.runs:
        style_run(run, size=10)

    symbols = ['@', '#', '€', '|', '\\', '[', '{', '°']
    table = doc.add_table(rows=5, cols=6)
    fixed_widths(table, [1.55, 2.25, 4.80, 1.55, 2.25, 4.80])

    headers = ['ZEICHEN', 'ALTGR +', '2. TASTE', 'ZEICHEN', 'ALTGR +', '2. TASTE']
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        margins(cell, 50, 45, 50, 45)
        border(cell, DARK, '12')
        shade(cell, DARK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = clear(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, bold=True, color=WHITE, size=10)

    for row_index in range(4):
        row = table.rows[row_index + 1]
        for side in range(2):
            item_index = row_index * 2 + side
            base = side * 3
            symbol_cell, combo_cell, answer_cell = row.cells[base:base + 3]

            for cell in (symbol_cell, combo_cell, answer_cell):
                margins(cell, 65, 60, 65, 60)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            border(symbol_cell, DARK, '12')
            shade(symbol_cell, ACCENT)
            p = clear(symbol_cell)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f'{item_index + 1}. {symbols[item_index]}')
            style_run(r, bold=True, color=WHITE, size=11)

            border(combo_cell, DARK, '12')
            shade(combo_cell, WHITE)
            p = clear(combo_cell)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run('AltGr +')
            style_run(r, bold=True, color=DARK, size=10)

            border(answer_cell, DARK, '20')
            shade(answer_cell, WHITE)
            p = clear(answer_cell)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        row.height = Cm(0.82)

    table.rows[0].height = Cm(0.62)
    instruction._p.addnext(table._tbl)
    old_table._tbl.getparent().remove(old_table._tbl)


def main():
    if len(sys.argv) != 2:
        raise SystemExit('Usage: python scripts/p1_altgr_readability.py <docx>')
    path = Path(sys.argv[1])
    if not path.exists():
        raise SystemExit(f'{path} not found')
    doc = Document(path)
    rebuild_altgr(doc)
    doc.save(path)
    print(f'{path}: rebuilt AltGr task for high-contrast readability')


if __name__ == '__main__':
    main()
