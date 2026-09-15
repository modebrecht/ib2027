from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn

PATH = Path('hw/P1.docx')
MAX_UNDERSCORES = 12
SAFETY_TWIPS = 120  # ~2.1 mm inside the document margins
DARK_BORDER_COLORS = {'0F172A', '334155'}
MIN_ANSWER_BORDER_SIZE = 8


def fail(message):
    raise SystemExit(f'P1 DOCX QA failed: {message}')


def all_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from all_paragraphs(cell)


def usable_width_twips(section):
    return int((section.page_width - section.left_margin - section.right_margin) / 635)


def table_grid_width_twips(table):
    total = 0
    for grid_col in table._tbl.tblGrid.gridCol_lst:
        value = grid_col.get(qn('w:w'))
        if value:
            total += int(value)
    return total


def paragraph_bottom_border(paragraph):
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        return None
    return p_bdr.find(qn('w:bottom'))


def cell_border(cell, edge_name):
    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return None
    borders = tc_pr.find(qn('w:tcBorders'))
    if borders is None:
        return None
    return borders.find(qn(f'w:{edge_name}'))


def assert_dark_cell_frame(cell, label, minimum_size=MIN_ANSWER_BORDER_SIZE):
    for edge_name in ('top', 'left', 'bottom', 'right'):
        edge = cell_border(cell, edge_name)
        if edge is None:
            fail(f'{label}: missing {edge_name} border')
        color = (edge.get(qn('w:color')) or '').upper()
        try:
            size = int(edge.get(qn('w:sz')) or '0')
        except ValueError:
            size = 0
        if color not in DARK_BORDER_COLORS:
            fail(f'{label}: low-contrast {edge_name} border color {color!r}')
        if size < minimum_size:
            fail(f'{label}: {edge_name} border too thin ({size} < {minimum_size})')


if not PATH.exists():
    fail(f'{PATH} does not exist')

doc = Document(PATH)
if not doc.sections:
    fail('document has no section')

usable = min(usable_width_twips(section) for section in doc.sections)
for idx, table in enumerate(doc.tables, 1):
    width = table_grid_width_twips(table)
    if width <= 0:
        fail(f'table {idx} has no explicit grid width')
    if width > usable - SAFETY_TWIPS:
        fail(f'table {idx} grid is {width} twips; usable width is {usable} twips')
    layout = table._tbl.tblPr.find(qn('w:tblLayout'))
    if layout is None or layout.get(qn('w:type')) != 'fixed':
        fail(f'table {idx} is not fixed-layout')

for p in all_paragraphs(doc):
    if re.search(r'_{%d,}' % (MAX_UNDERSCORES + 1), p.text):
        fail(f'long unbreakable underscore sequence remains: {p.text[:50]!r}')
    if re.fullmatch(r'_{4,}', p.text.strip()):
        fail('underscore-only answer placeholder remains; field is not direct-write')

    # Writable answer lines are real paragraph borders after the layout pass.
    # They must be dark and thick enough to remain visible on old monitors and
    # washed-out projectors; pale grey is not acceptable as a structural cue.
    bottom = paragraph_bottom_border(p)
    if bottom is not None:
        color = (bottom.get(qn('w:color')) or '').upper()
        try:
            size = int(bottom.get(qn('w:sz')) or '0')
        except ValueError:
            size = 0
        if color not in DARK_BORDER_COLORS:
            fail(f'answer line uses low-contrast border color {color!r}')
        if size < MIN_ANSWER_BORDER_SIZE:
            fail(f'answer line border too thin ({size} < {MIN_ANSWER_BORDER_SIZE})')

if any(p._p.xpath('.//w:br[@w:type="page"]') for p in doc.paragraphs):
    fail('literal page-break paragraph remains and may create a blank page after reflow')

green_heading = next((p for p in doc.paragraphs if p.text.strip().startswith('8. Green IT')), None)
if green_heading is None:
    fail('Green IT heading missing')
if green_heading.paragraph_format.page_break_before:
    fail('Green IT is still forced onto a new page')

if not any('Bonus 1 – Hardware Detective' in p.text for p in doc.paragraphs):
    fail('bonus section missing')

if not any('Nenne ein passendes Beispiel für die Eingabe dieser EVA-Kette' in p.text for p in doc.paragraphs):
    fail('direct-write EVA input question missing')

# Regression checks for the two dense matching tables: the response areas
# must stay explicit dark boxes rather than white/very-light-grey regions.
eva_table = next((
    table for table in doc.tables
    if table.rows and table.rows[0].cells[0].text.strip() == 'Situation'
    and 'Antwort (E / V / A)' in table.rows[0].cells[1].text
), None)
if eva_table is None:
    fail('EVA table missing')
for row_index, row in enumerate(eva_table.rows[1:], 1):
    assert_dark_cell_frame(row.cells[1], f'EVA answer cell {row_index}', minimum_size=12)

ports_table = next((
    table for table in doc.tables
    if any('Moderner Monitor oder Fernseher' in cell.text for row in table.rows for cell in row.cells)
    and any('Interne SSD direkt auf dem Mainboard' in cell.text for row in table.rows for cell in row.cells)
), None)
if ports_table is None:
    fail('Anschlüsse table missing')
for row_index, row in enumerate(ports_table.rows, 1):
    assert_dark_cell_frame(row.cells[1], f'Anschlüsse answer cell {row_index}', minimum_size=12)

print(f'P1 DOCX QA passed: {len(doc.tables)} tables, direct-write fields, high-contrast answer areas, compact natural page flow.')
